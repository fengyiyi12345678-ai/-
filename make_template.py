"""生成 Excel 数据收集模板：竞品总表 + 每个商品一张评价表。"""
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUTPUT = Path(__file__).parent / "competitor_template.xlsx"
PRODUCT_COUNT = 20

HEADER_FILL = PatternFill("solid", fgColor="2E75B6")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HINT_FILL = PatternFill("solid", fgColor="FFF2CC")


def style_header(ws, row, cols):
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def build_summary(ws):
    ws.title = "竞品总表"
    headers = [
        "序号", "商品标题", "店铺名", "商品链接",
        "价格(元)", "月销量(件)", "累计评价数", "好评率(%)",
        "主图特点(自填)", "备注",
    ]
    ws.append(headers)
    style_header(ws, 1, len(headers))

    hint = ws.cell(row=2, column=1, value="↓ 在淘宝按关键词搜索后，按销量排序，前 20 个依次填入")
    hint.fill = HINT_FILL
    hint.font = Font(italic=True, color="7F6000")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))

    for i in range(1, PRODUCT_COUNT + 1):
        ws.cell(row=2 + i, column=1, value=i)

    widths = [6, 40, 18, 30, 10, 12, 12, 10, 25, 20]
    for idx, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    ws.freeze_panes = "A3"


def build_review_sheet(wb, idx):
    ws = wb.create_sheet(f"P{idx:02d}_评价")
    headers = ["评价类型", "评价内容", "用户(可选)", "日期(可选)", "SKU(可选)"]
    ws.append([f"商品 {idx} 的评价 — 把淘宝评价区文字复制到 B 列即可"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.cell(row=1, column=1).fill = HINT_FILL
    ws.cell(row=1, column=1).font = Font(italic=True, color="7F6000")
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="center")

    ws.append(headers)
    style_header(ws, 2, len(headers))

    type_hint = (
        '评价类型留空即可（脚本自动判断），'
        '或手动写"好评/中评/差评"覆盖判断'
    )
    ws.cell(row=3, column=1, value=type_hint).font = Font(italic=True, color="999999")

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 15
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 20
    ws.freeze_panes = "A3"


def main():
    wb = openpyxl.Workbook()
    build_summary(wb.active)
    for i in range(1, PRODUCT_COUNT + 1):
        build_review_sheet(wb, i)
    wb.save(OUTPUT)
    print(f"模板已生成: {OUTPUT}")
    print(f"  - 1 张总表 + {PRODUCT_COUNT} 张评价表")


if __name__ == "__main__":
    main()
